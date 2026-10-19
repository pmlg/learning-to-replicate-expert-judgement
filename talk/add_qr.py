"""DAT-340: generate repo QR code and add it to the deck's final slide."""
import qrcode
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

REPO_URL = "https://github.com/deep-science-ai/qwen-chess-rl"
TALK_DIR = "/home/ibrahim/workspace/qwen-chess-rl/talk"
PPTX = f"{TALK_DIR}/pmlg_chess_rl_talk_draft.pptx"
QR = f"{TALK_DIR}/repo_qr.png"

# 1. QR code
img = qrcode.make(REPO_URL, box_size=10, border=2)
img.save(QR)
print("QR saved:", QR, img.size)

# 2. Add to final slide (26)
prs = Presentation(PPTX)
slide = prs.slides[len(prs.slides) - 1]
assert "Thank you" in slide.shapes.title.text, slide.shapes.title.text

# remove any pre-existing QR (idempotent re-run)
for sh in list(slide.shapes):
    if sh.name in ("DAT340_REPO_QR", "DAT340_REPO_QR_CAPTION"):
        sh._element.getparent().remove(sh._element)

qr = slide.shapes.add_picture(QR, Inches(10.35), Inches(2.05), height=Inches(2.4))
qr.name = "DAT340_REPO_QR"

cap = slide.shapes.add_textbox(Inches(9.95), Inches(4.55), Inches(3.15), Inches(1.3))
cap.name = "DAT340_REPO_QR_CAPTION"
tf = cap.text_frame
tf.word_wrap = True
lines = ["Scan for slides, code & the", "step-by-step notebook:", "github.com/deep-science-ai/", "qwen-chess-rl"]
for j, line in enumerate(lines):
    para = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
    run = para.add_run()
    run.text = line
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(0xC0, 0xC0, 0xC0)
    run.font.name = "Consolas" if line.startswith(("github", "qwen")) else "Calibri"

prs.save(PPTX)
print("saved deck")

# 3. verify
prs2 = Presentation(PPTX)
s = prs2.slides[len(prs2.slides) - 1]
pics = [sh for sh in s.shapes if sh.shape_type == 13]
names = [sh.name for sh in s.shapes]
assert "DAT340_REPO_QR" in names, names
print("verify: slides =", len(prs2.slides), "| final-slide shapes =", names, "| pictures =", len(pics))
